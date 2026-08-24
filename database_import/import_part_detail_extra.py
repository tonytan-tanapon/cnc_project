#!/usr/bin/env python3

from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

import sys
import os


# =====================================================
# CONFIG
# =====================================================

DATABASE_URL = (
    "postgresql+psycopg2://postgres:1234@localhost:5432/mydb"
)

DATA_FOLDER = Path(
    r"Z:\Topnotch Group\Public\Data Base & Inventory Stock\Data"
)

# True  = ทดสอบอย่างเดียว ไม่ update database
# False = update database จริง
PREVIEW_MODE = False


# =====================================================
# IMPORT MODELS
# =====================================================

sys.path.append(
    os.path.abspath(
        os.path.join(
            os.path.dirname(__file__),
            ".."
        )
    )
)

from models import Part


# =====================================================
# DATABASE
# =====================================================

engine = create_engine(
    DATABASE_URL,
    future=True
)

SessionLocal = sessionmaker(
    bind=engine,
    class_=Session,
    autoflush=False,
    autocommit=False,
)


# =====================================================
# EXTRACT DATA FROM EXCEL
# =====================================================

def extract_part_detail_extra(file_path: Path):

    try:
        # ชื่อไฟล์ = Part No = ชื่อ Sheet
        part_no = file_path.stem.strip()

        wb = load_workbook(
            file_path,
            data_only=True,
            read_only=True
        )

        # ==========================================
        # หา Sheet
        # ==========================================

        if part_no not in wb.sheetnames:
            print(f"    SHEET NOT FOUND: {part_no}")
            print(f"    Available sheets: {wb.sheetnames}")
            return None

        ws = wb[part_no]

        print(f"    Sheet: {ws.title}")

        # ==========================================
        # STEP 1
        # อ่าน Column B เริ่ม Row 7
        # ==========================================

        row = 7

        while True:

            value_b = ws.cell(
                row=row,
                column=2       # Column B
            ).value

            # ถ้า B ไม่มีข้อมูล
            # ให้เริ่มอ่าน Column J จาก row นี้
            if value_b is None or str(value_b).strip() == "":
                break

            row += 1

        # row ตอนนี้คือแถวแรกที่ Column B ว่าง
        start_row_j = row

        print(
            f"    Column B ended at row: {row - 1}"
        )

        print(
            f"    Start reading Column J at row: {start_row_j}"
        )

        # ==========================================
        # STEP 2
        # อ่าน Column I และ J
        # ==========================================

        results = []

        empty_count = 0
        row = start_row_j

        while True:

            value_i = ws.cell(
                row=row,
                column=9       # Column I
            ).value

            value_j = ws.cell(
                row=row,
                column=10      # Column J
            ).value

            # clean ค่า
            text_i = "" if value_i is None else str(value_i).strip()
            text_j = "" if value_j is None else str(value_j).strip()

            # ======================================
            # ถ้า I และ J ว่างทั้งคู่
            # ======================================

            if not text_i and not text_j:

                empty_count += 1

                print(
                    f"    I{row}, J{row}: EMPTY "
                    f"({empty_count}/5)"
                )

                # ว่าง 5 แถวติดกัน
                if empty_count >= 5:
                    print(
                        "    5 consecutive empty rows -> STOP"
                    )
                    break

            else:

                # เจอข้อมูล reset ตัวนับ
                empty_count = 0

                # ==================================
                # รวมข้อมูล I + J ของแถวนั้น
                # ==================================

                row_values = []

                if text_i:
                    row_values.append(text_i)

                if text_j:
                    row_values.append(text_j)

                row_text = " ".join(row_values)

                results.append(f"• {row_text}")

                print(
                    f"    Row {row}: {row_text}"
                )

            row += 1

        # ==========================================
        # STEP 3
        # ไม่มีข้อมูล
        # ==========================================

        if not results:
            print(
                "    No data found in Column J"
            )
            return None

        # ==========================================
        # STEP 4
        # รวมข้อมูล
        # ==========================================

        extra_detail = "\n".join(results)

        print("    --------------------")
        print("    Extracted:")
        print(extra_detail)
        print("    --------------------")

        return extra_detail

    except Exception as e:

        print(
            f"    ERROR reading Excel: {e}"
        )

        return None


# =====================================================
# GET PART NUMBER FROM FILE NAME
# =====================================================

def get_part_no_from_filename(
    file_path: Path
) -> str:
    """
    ตัวอย่าง:

    ABC-001.xlsx
        -> ABC-001

    123-456.xlsm
        -> 123-456
    """

    return file_path.stem.strip()


# =====================================================
# MAIN
# =====================================================

def main():

    print("=" * 60)
    print("IMPORT PART DETAIL EXTRA")
    print("=" * 60)

    print(f"Folder : {DATA_FOLDER}")
    print(f"Preview: {PREVIEW_MODE}")
    print()

    # -----------------------------------------------
    # ตรวจสอบว่า Folder มีจริงหรือไม่
    # -----------------------------------------------

    if not DATA_FOLDER.exists():

        print(
            f"ERROR: Folder not found:\n"
            f"{DATA_FOLDER}"
        )

        return

    if not DATA_FOLDER.is_dir():

        print(
            f"ERROR: Path is not a folder:\n"
            f"{DATA_FOLDER}"
        )

        return


    # -----------------------------------------------
    # Counters
    # -----------------------------------------------

    total_files = 0

    part_found = 0
    part_not_found = 0

    data_found = 0
    data_not_found = 0

    updated = 0
    errors = 0


    # -----------------------------------------------
    # Database Session
    # -----------------------------------------------

    with SessionLocal() as db:

        # ===========================================
        # Loop files
        # ===========================================

        for file_path in DATA_FOLDER.iterdir():

            # ---------------------------------------
            # ข้าม folder
            # ---------------------------------------

            if not file_path.is_file():
                continue


            # ---------------------------------------
            # เอาเฉพาะ Excel
            # ---------------------------------------

            if file_path.suffix.lower() not in (
                ".xlsx",
                ".xlsm",
            ):
                continue


            # ---------------------------------------
            # ข้าม temporary Excel files
            #
            # เช่น:
            # ~$ABC-001.xlsx
            # ---------------------------------------

            if file_path.name.startswith("~$"):
                continue


            total_files += 1

            print()
            print("-" * 60)

            print(
                f"[{total_files}] "
                f"{file_path.name}"
            )


            # =======================================
            # PART NUMBER FROM FILE NAME
            # =======================================

            part_no = get_part_no_from_filename(file_path)

            print(f"Part No: {part_no}")


            # =======================================
            # FIND PART IN DATABASE
            # =======================================

            try:

                part = db.execute(
                    select(Part).where(
                        Part.part_no == part_no
                    )
                ).scalar_one_or_none()

            except Exception as e:

                print(f"    DATABASE ERROR: {e}")
                errors += 1
                continue


            # ---------------------------------------
            # Part ไม่อยู่ใน DB
            # ---------------------------------------

            if not part:

                print(f"    NOT FOUND IN DATABASE")
                part_not_found += 1
                continue


            part_found += 1
            print(f"    Part ID: {part.id}")


            # =======================================
            # READ EXCEL
            # =======================================

            try:

                extra_detail = (
                    extract_part_detail_extra(
                        file_path
                    )
                )

            except Exception as e:

                print(f"    EXTRACT ERROR: {e}")
                errors += 1
                continue


            # =======================================
            # NO DATA
            # =======================================

            if extra_detail is None:

                print("    No extra detail found")
                data_not_found += 1
                continue


            # =======================================
            # CLEAN DATA
            # =======================================

            extra_detail = str(
                extra_detail
            ).strip()


            if not extra_detail:

                print("    Extra detail is empty")

                data_not_found += 1

                continue


            data_found += 1


            # =======================================
            # SHOW OLD / NEW VALUE
            # =======================================

            old_value = (
                part.part_detail_extra
                if part.part_detail_extra
                else ""
            )

            print(f"    OLD: {old_value}")

            print(f"    NEW: {extra_detail}")


            # =======================================
            # UPDATE
            # =======================================

            if PREVIEW_MODE:

                print("    PREVIEW ONLY - NOT UPDATED")

            else:

                part.part_detail_extra = (extra_detail)

                updated += 1

                print("    UPDATED")


        # ===========================================
        # COMMIT / ROLLBACK
        # ===========================================

        if PREVIEW_MODE:

            db.rollback()

            print()
            print("PREVIEW MODE - DATABASE NOT CHANGED")

        else:

            try:

                db.commit()

                print()
                print(
                    "DATABASE COMMITTED"
                )

            except Exception as e:

                db.rollback()

                print()
                print(
                    f"COMMIT ERROR: {e}"
                )

                return


    # =================================================
    # SUMMARY
    # =================================================

    print()
    print("=" * 60)
    print("SUMMARY")
    print("=" * 60)

    print(        f"Excel files     : {total_files}"    )

    print(        f"Parts found     : {part_found}"    )

    print(        f"Parts not found : {part_not_found}"    )

    print(        f"Data found      : {data_found}"    )

    print(        f"Data not found  : {data_not_found}"    )

    print(        f"Updated         : {updated}"    )

    print(        f"Errors          : {errors}"    )

    print("=" * 60)


# =====================================================
# RUN
# =====================================================

if __name__ == "__main__":
    main()
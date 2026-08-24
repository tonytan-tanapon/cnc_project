#!/usr/bin/env python3

from __future__ import annotations

from openpyxl import load_workbook
from pathlib import Path
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

import sys
import os


# ---------- CONFIG ----------

DATABASE_URL = "postgresql+psycopg2://postgres:1234@localhost:5432/mydb"

XLSX_FILE = Path(
    r"Z:\Topnotch Group\Public\Data Base & Inventory Stock\Data Base.xlsm"
)


# ---------- IMPORT MODELS ----------

sys.path.append(
    os.path.abspath(
        os.path.join(os.path.dirname(__file__), "..")
    )
)

from models import Part


# ---------- ENGINE ----------

engine = create_engine(DATABASE_URL, future=True)

SessionLocal = sessionmaker(
    bind=engine,
    class_=Session,
    autoflush=False,
    autocommit=False,
)


# =====================================================
# MAIN
# =====================================================

def main():

    with SessionLocal() as db:

        wb = load_workbook(
            XLSX_FILE,
            data_only=True,
            read_only=True
        )

        # เปลี่ยนชื่อ sheet ตรงนี้
        # ws = wb["Sheet1"]

        ws = wb.worksheets[0]

        # -----------------------------
        # หา column จาก header
        # -----------------------------
        # Header อยู่แถว 4
        headers = {
            str(cell.value).strip(): cell.column
            for cell in ws[4]
            if cell.value is not None
        }

        print("Headers:", headers)

        part_col = headers["Part Number"]
        detail_col = headers["FAIR/LONG FORM/NOTE"]

        # ตัวนับ
        updated = 0
        not_found = 0

        # ข้อมูลเริ่มแถว 5
        for row in range(5, ws.max_row + 1):

            part_no = ws.cell(row, part_col).value
            part_detail = ws.cell(row, detail_col).value

            if not part_no:
                continue

            part_no = str(part_no).strip()

            if part_detail is None:
                continue

            part_detail = str(part_detail).strip()

            part = db.execute(
                select(Part).where(
                    Part.part_no == part_no
                )
            ).scalar_one_or_none()

            if not part:
                print(f"NOT FOUND: {part_no}")
                not_found += 1
                continue

            part.part_detail = part_detail
            updated += 1

            print(f"UPDATE: {part_no} -> {part_detail}")

        db.commit()

        print("-----------------------------")
        print(f"Updated   : {updated}")
        print(f"Not found : {not_found}")
        print("-----------------------------")


if __name__ == "__main__":
    main()